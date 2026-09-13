"""
Report module: turn a summarized orders DataFrame into a formatted
Excel report with a styled table, a totals row, and a bar chart.

This module owns all file-writing/formatting concerns. It expects the
DataFrame produced by processing.summarize.summarize_orders, i.e. one
row per customer with columns:
    customer, total_amount, order_count, average_order, last_order_date
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTALS_FONT = Font(bold=True)
TOTALS_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

COLUMN_HEADERS = {
    "customer": "Customer",
    "total_amount": "Total Amount",
    "order_count": "Order Count",
    "average_order": "Average Order",
    "last_order_date": "Last Order Date",
}
CURRENCY_COLUMNS = {"total_amount", "average_order"}
DATE_COLUMNS = {"last_order_date"}
CURRENCY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"


def generate_report(
    summary: pd.DataFrame,
    output_path: str | Path,
    sheet_name: str = "Sales Report",
) -> Path:
    """
    Write a formatted Excel report from a per-customer summary DataFrame.

    Parameters
    ----------
    summary : pd.DataFrame
        Output of summarize_orders: one row per customer with columns
        customer, total_amount, order_count, average_order,
        last_order_date.
    output_path : str | Path
        Where to save the .xlsx file. Parent directories are created
        if they don't exist.
    sheet_name : str, default "Sales Report"
        Name of the worksheet containing the table and chart.

    Returns
    -------
    Path
        The path the report was written to.

    Raises
    ------
    KeyError
        If summary is missing any of the expected columns.
    """
    expected_columns = list(COLUMN_HEADERS.keys())
    missing = [col for col in expected_columns if col not in summary.columns]
    if missing:
        raise KeyError(f"Missing expected column(s) in summary: {missing}")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    _write_header_row(sheet, expected_columns)
    _write_data_rows(sheet, summary, expected_columns)
    totals_row_index = _write_totals_row(sheet, summary, expected_columns)
    _autofit_columns(sheet, summary, expected_columns)
    _add_bar_chart(sheet, summary, totals_row_index)

    workbook.save(output_path)
    return output_path


def _write_header_row(sheet: Worksheet, columns: list[str]) -> None:
    for col_index, column_key in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_index, value=COLUMN_HEADERS[column_key])
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _write_data_rows(
    sheet: Worksheet, summary: pd.DataFrame, columns: list[str]
) -> None:
    for row_offset, (_, record) in enumerate(summary.iterrows(), start=2):
        for col_index, column_key in enumerate(columns, start=1):
            value = record[column_key]
            cell = sheet.cell(row=row_offset, column=col_index, value=value)
            if column_key in CURRENCY_COLUMNS:
                cell.number_format = CURRENCY_FORMAT
            elif column_key in DATE_COLUMNS and pd.notna(value):
                cell.number_format = DATE_FORMAT


def _write_totals_row(
    sheet: Worksheet, summary: pd.DataFrame, columns: list[str]
) -> int:
    totals_row_index = len(summary) + 2  # +1 for header, +1 for 1-based index

    for col_index, column_key in enumerate(columns, start=1):
        if column_key == "customer":
            value = "Total"
        elif column_key == "total_amount":
            value = summary["total_amount"].sum()
        elif column_key == "order_count":
            value = summary["order_count"].sum()
        elif column_key == "average_order":
            total = summary["total_amount"].sum()
            count = summary["order_count"].sum()
            value = (total / count) if count else 0
        else:
            value = None

        cell = sheet.cell(row=totals_row_index, column=col_index, value=value)
        cell.font = TOTALS_FONT
        cell.fill = TOTALS_FILL
        if column_key in CURRENCY_COLUMNS:
            cell.number_format = CURRENCY_FORMAT

    return totals_row_index


def _autofit_columns(
    sheet: Worksheet, summary: pd.DataFrame, columns: list[str]
) -> None:
    for col_index, column_key in enumerate(columns, start=1):
        header_len = len(COLUMN_HEADERS[column_key])
        if len(summary) > 0:
            max_value_len = summary[column_key].astype(str).map(len).max()
        else:
            max_value_len = 0
        width = max(header_len, max_value_len) + 4
        sheet.column_dimensions[get_column_letter(col_index)].width = width


def _add_bar_chart(
    sheet: Worksheet, summary: pd.DataFrame, totals_row_index: int
) -> None:
    if summary.empty:
        return

    chart = BarChart()
    chart.title = "Total Amount by Customer"
    chart.x_axis.title = "Customer"
    chart.y_axis.title = "Total Amount"
    chart.style = 10

    last_data_row = totals_row_index - 1  # exclude the totals row from the chart

    data = Reference(
        sheet, min_col=2, max_col=2, min_row=1, max_row=last_data_row
    )  # column 2 = total_amount, includes header for series title
    categories = Reference(
        sheet, min_col=1, max_col=1, min_row=2, max_row=last_data_row
    )  # column 1 = customer names

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 20
    chart.height = 10

    anchor_column = get_column_letter(len(COLUMN_HEADERS) + 2)
    sheet.add_chart(chart, f"{anchor_column}2")