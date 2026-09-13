"""
Unit tests for report.excel_report.generate_report.

Run with:
    pytest tests/test_excel_report.py -v
"""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from report.excel_report import generate_report


@pytest.fixture
def summary_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "customer": ["Bruno", "Ana", "Carla"],
            "total_amount": [500.0, 150.0, 75.0],
            "order_count": [2, 2, 1],
            "average_order": [250.0, 75.0, 75.0],
            "last_order_date": pd.to_datetime(
                ["2026-03-01", "2026-02-10", "2026-01-20"]
            ),
        }
    )


def test_creates_file_at_output_path(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"

    result_path = generate_report(summary_df, output_path)

    assert result_path == output_path
    assert output_path.exists()


def test_creates_parent_directories(tmp_path: Path, summary_df):
    output_path = tmp_path / "nested" / "dir" / "report.xlsx"

    generate_report(summary_df, output_path)

    assert output_path.exists()


def test_header_row_has_expected_labels(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]

    header_values = [cell.value for cell in sheet[1]]
    assert header_values == [
        "Customer",
        "Total Amount",
        "Order Count",
        "Average Order",
        "Last Order Date",
    ]


def test_data_rows_match_input(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]

    # Row 2 corresponds to the first row of summary_df (Bruno).
    assert sheet.cell(row=2, column=1).value == "Bruno"
    assert sheet.cell(row=2, column=2).value == 500.0
    assert sheet.cell(row=2, column=3).value == 2


def test_totals_row_sums_amount_and_count(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]

    totals_row = len(summary_df) + 2  # header + 3 data rows -> row 5
    assert sheet.cell(row=totals_row, column=1).value == "Total"
    assert sheet.cell(row=totals_row, column=2).value == 725.0  # 500+150+75
    assert sheet.cell(row=totals_row, column=3).value == 5  # 2+2+1


def test_currency_columns_have_number_format(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]

    # total_amount is column 2.
    assert sheet.cell(row=2, column=2).number_format == "#,##0.00"


def test_chart_is_added_to_sheet(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]

    assert len(sheet._charts) == 1


def test_missing_expected_column_raises_key_error(tmp_path: Path):
    incomplete_df = pd.DataFrame({"customer": ["Ana"], "total_amount": [100.0]})
    output_path = tmp_path / "report.xlsx"

    with pytest.raises(KeyError):
        generate_report(incomplete_df, output_path)


def test_empty_summary_still_generates_file_without_chart(tmp_path: Path):
    empty_df = pd.DataFrame(
        columns=[
            "customer",
            "total_amount",
            "order_count",
            "average_order",
            "last_order_date",
        ]
    )
    output_path = tmp_path / "report.xlsx"

    generate_report(empty_df, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["Sales Report"]
    assert len(sheet._charts) == 0
    # Totals row should still be written (row 2, since there are 0 data rows).
    assert sheet.cell(row=2, column=1).value == "Total"


def test_custom_sheet_name(tmp_path: Path, summary_df):
    output_path = tmp_path / "report.xlsx"
    generate_report(summary_df, output_path, sheet_name="Q1 Summary")

    workbook = load_workbook(output_path)
    assert "Q1 Summary" in workbook.sheetnames