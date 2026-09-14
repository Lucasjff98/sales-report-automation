"""
Integration tests for main.py: exercise the full pipeline end-to-end,
using real files on disk but a mocked SMTP server (no real network
calls or credentials required).

Run with:
    pytest tests/test_main.py -v
"""

from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from main import parse_column_map, run_pipeline


@pytest.fixture
def raw_orders_csv(tmp_path: Path) -> Path:
    df = pd.DataFrame(
        {
            "Cliente": ["Ana", "Bruno", "Ana"],
            "Valor": [100.0, 200.0, 50.0],
            "Data": ["2026-01-01", "2026-01-05", "2026-02-10"],
        }
    )
    csv_path = tmp_path / "raw_orders.csv"
    df.to_csv(csv_path, index=False)
    return csv_path


@pytest.fixture
def delivery_config_file(tmp_path: Path) -> Path:
    config_path = tmp_path / "config.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "smtp_host": "smtp.office365.com",
                "smtp_port": 587,
                "recipients": ["boss@example.com"],
            }
        )
    )
    return config_path


@pytest.fixture
def env_file(tmp_path: Path) -> Path:
    env_path = tmp_path / ".env"
    env_path.write_text(
        "SENDER_EMAIL=reports@example.com\nSENDER_PASSWORD=app-password-123\n"
    )
    return env_path


def make_args(
    tmp_path: Path,
    raw_orders_csv: Path,
    delivery_config_file: Path,
    env_file: Path,
    **overrides,
):
    from argparse import Namespace

    defaults = dict(
        input=str(raw_orders_csv),
        output=str(tmp_path / "output" / "report.xlsx"),
        column_map="Cliente=customer,Valor=amount,Data=order_date",
        sheet_name=0,
        config=str(delivery_config_file),
        env_file=str(env_file),
        subject=None,
        no_email=False,
    )
    defaults.update(overrides)
    return Namespace(**defaults)


# ---------------------------------------------------------------------------
# parse_column_map
# ---------------------------------------------------------------------------

def test_parse_column_map_returns_none_for_empty_input():
    assert parse_column_map(None) is None
    assert parse_column_map("") is None


def test_parse_column_map_parses_multiple_pairs():
    result = parse_column_map("Cliente=customer,Valor=amount,Data=order_date")

    assert result == {
        "Cliente": "customer",
        "Valor": "amount",
        "Data": "order_date",
    }


def test_parse_column_map_raises_on_invalid_format():
    with pytest.raises(ValueError, match="Invalid --column-map entry"):
        parse_column_map("Cliente-customer")


# ---------------------------------------------------------------------------
# run_pipeline (full integration, email mocked)
# ---------------------------------------------------------------------------

@patch("main.send_report_email")
def test_full_pipeline_generates_report_and_sends_email(
    mock_send_email, tmp_path, raw_orders_csv, delivery_config_file, env_file
):
    args = make_args(tmp_path, raw_orders_csv, delivery_config_file, env_file)

    report_path = run_pipeline(args)

    assert Path(report_path).exists()
    mock_send_email.assert_called_once()


@patch("main.send_report_email")
def test_report_content_reflects_aggregated_orders(
    mock_send_email, tmp_path, raw_orders_csv, delivery_config_file, env_file
):
    args = make_args(tmp_path, raw_orders_csv, delivery_config_file, env_file)

    report_path = run_pipeline(args)

    workbook = load_workbook(report_path)
    sheet = workbook["Sales Report"]
    # Ana had 2 orders totaling 150.0, Bruno had 1 order of 200.0 ->
    # sorted descending by total_amount, Bruno should be first.
    assert sheet.cell(row=2, column=1).value == "Bruno"
    assert sheet.cell(row=2, column=2).value == 200.0
    assert sheet.cell(row=3, column=1).value == "Ana"
    assert sheet.cell(row=3, column=2).value == 150.0


@patch("main.send_report_email")
def test_no_email_flag_skips_sending(
    mock_send_email, tmp_path, raw_orders_csv, delivery_config_file, env_file
):
    args = make_args(
        tmp_path, raw_orders_csv, delivery_config_file, env_file, no_email=True
    )

    report_path = run_pipeline(args)

    assert Path(report_path).exists()
    mock_send_email.assert_not_called()


@patch("main.send_report_email")
def test_custom_subject_is_forwarded_to_send_email(
    mock_send_email, tmp_path, raw_orders_csv, delivery_config_file, env_file
):
    args = make_args(
        tmp_path,
        raw_orders_csv,
        delivery_config_file,
        env_file,
        subject="Custom Subject",
    )

    run_pipeline(args)

    _, kwargs = mock_send_email.call_args
    assert kwargs["subject"] == "Custom Subject"


def test_pipeline_raises_when_input_file_missing(
    tmp_path, delivery_config_file, env_file
):
    args = make_args(
        tmp_path,
        tmp_path / "does_not_exist.csv",
        delivery_config_file,
        env_file,
    )

    with pytest.raises(FileNotFoundError):
        run_pipeline(args)